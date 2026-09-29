import os
import shutil
from pathlib import Path
from urllib.parse import urlparse

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware

from backend.vector_store import add_document, search

app = FastAPI(
    title="iSmartRAG API",
    description="Backend API for iSmartRAG",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
allow_origins=[
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    "http://localhost:5174",
    "http://127.0.0.1:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = Path(__file__).resolve().parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)


@app.get("/")
def root():
    return {
        "message": "iSmartRAG API is running"
    }


@app.get("/health")
def health():
    return {
        "status": "healthy"
    }


def parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = []

    for index, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        text.append(f"PAGE {index + 1}\n{page_text}")

    return "\n\n".join(text)


def parse_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    for table_index, table in enumerate(document.tables):
        lines.append(f"\nTABLE {table_index + 1}")

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))

    return "\n".join(lines)


def parse_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    lines = []

    for slide_index, slide in enumerate(presentation.slides):
        lines.append(f"\nSLIDE {slide_index + 1}")

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text)

    return "\n".join(lines)


def parse_excel(path: Path) -> str:
    import pandas as pd

    suffix = path.suffix.lower()

    if suffix == ".csv":
        sheets = {
            "CSV": pd.read_csv(path)
        }
    else:
        sheets = pd.read_excel(path, sheet_name=None)

    lines = []

    for sheet_name, dataframe in sheets.items():
        lines.append(f"\nSHEET: {sheet_name}")
        lines.append(f"ROWS: {len(dataframe)}")
        lines.append(f"COLUMNS: {len(dataframe.columns)}")

        lines.append(
            dataframe.fillna("").to_string(index=False)
        )

    return "\n".join(lines)


def parse_file(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return parse_pdf(path), "pdf"

    if suffix in {".docx", ".doc"}:
        if suffix == ".doc":
            raise ValueError(
                "Legacy .doc files are not supported. Please upload .docx."
            )

        return parse_docx(path), "docx"

    if suffix in {".ppt", ".pptx"}:
        if suffix == ".ppt":
            raise ValueError(
                "Legacy .ppt files are not supported. Please upload .pptx."
            )

        return parse_pptx(path), "ppt"

    if suffix in {
        ".xlsx",
        ".xlsm",
        ".xltx",
        ".xls",
        ".csv",
    }:
        return parse_excel(path), "excel"

    raise ValueError(
        f"Unsupported file type: {suffix}"
    )


@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No filename provided.",
        )

    filename = Path(file.filename).name
    file_path = UPLOAD_DIR / filename

    try:
        with file_path.open("wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        text, doc_type = parse_file(file_path)

        if not text.strip():
            raise HTTPException(
                status_code=400,
                detail="No readable text was found in the file.",
            )

        chunk_count = add_document(
            text,
            filename,
            doc_type,
        )

        return {
            "success": True,
            "filename": filename,
            "type": doc_type,
            "chunks": chunk_count,
            "message": "File uploaded and added to the knowledge base.",
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to process file: {exc}",
        )

    finally:
        file.file.close()


@app.post("/website")
def add_website(url: str):
    url = url.strip()

    if not url:
        raise HTTPException(
            status_code=400,
            detail="URL is required.",
        )

    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    parsed = urlparse(url)

    if not parsed.netloc:
        raise HTTPException(
            status_code=400,
            detail="Invalid URL.",
        )

    try:
        from playwright.sync_api import sync_playwright
        from bs4 import BeautifulSoup

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(
                headless=True
            )

            page = browser.new_page()

            page.goto(
                url,
                timeout=60000,
                wait_until="networkidle",
            )

            html = page.content()

            browser.close()

        soup = BeautifulSoup(html, "html.parser")

        for element in soup(
            ["script", "style", "noscript"]
        ):
            element.decompose()

        text = soup.get_text(
            separator="\n",
            strip=True,
        )

        if not text.strip():
            raise HTTPException(
                status_code=400,
                detail="No readable text was found on the website.",
            )

        chunk_count = add_document(
            text,
            url,
            "web",
        )

        return {
            "success": True,
            "url": url,
            "type": "web",
            "chunks": chunk_count,
            "message": "Website added to the knowledge base.",
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to scrape website: {exc}",
        )


@app.get("/search")
def search_documents(
    query: str,
    n_results: int = 3,
):
    query = query.strip()

    if not query:
        raise HTTPException(
            status_code=400,
            detail="Query is required.",
        )

    n_results = max(1, min(n_results, 10))

    try:
        results = search(
            query,
            n_results=n_results,
        )

        documents = results.get("documents", [[]])[0]
        metadatas = results.get("metadatas", [[]])[0]
        distances = results.get("distances", [[]])[0]

        formatted_results = []

        for index, document in enumerate(documents):
            metadata = (
                metadatas[index]
                if index < len(metadatas)
                else {}
            )

            distance = (
                distances[index]
                if index < len(distances)
                else None
            )

            formatted_results.append(
                {
                    "text": document,
                    "source": metadata.get("source"),
                    "type": metadata.get("type"),
                    "distance": distance,
                }
            )

        return {
            "query": query,
            "results": formatted_results,
        }

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Search failed: {exc}",
        )
